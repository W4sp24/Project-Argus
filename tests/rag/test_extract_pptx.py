"""Tests for PowerPoint extraction.

``python-pptx``'s ``slide.shapes`` is a shallow iterator over recognised
shape elements. Four kinds of content are invisible to it, and lecture decks
are full of all four: anything inside ``mc:AlternateContent`` (where OMML
equations live), anything inside a group, tables, and speaker notes.

Measured against the author's vault before this module existed, the two
maths-heavy decks in CS26110 reached the model with 29% and 32% of their text
and none of their 198 equations -- which is why a guide generated from them
said "the supplied text does not contain its definition" fourteen times, and
why a practice exam asked which theorem was "listed in the excerpt" instead of
asking anything about the theorem.
"""

from pathlib import Path

import pytest
from lxml import etree

from backend.rag.extract import extract_blocks
from backend.rag.extractors.pptx_shapes import slide_lines, speaker_notes

pytest.importorskip("pptx", reason="the [rag] extra is not installed")

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _deck(tmp_path: Path):
    """A one-slide presentation with a title, ready to be added to."""
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Limits"
    return presentation, slide


def _save(presentation, tmp_path: Path) -> Path:
    path = tmp_path / "deck.pptx"
    presentation.save(path)
    return path


def test_a_plain_slide_still_extracts(tmp_path: Path) -> None:
    """The behaviour that already worked must keep working."""
    presentation, slide = _deck(tmp_path)
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "A limit is a value a function approaches."
    assert slide_lines(slide) == [
        "Limits",
        "A limit is a value a function approaches.",
    ]


def test_text_inside_a_group_is_recovered(tmp_path: Path) -> None:
    """``slide.shapes`` does not descend into a group, so grouped text vanished."""
    presentation, slide = _deck(tmp_path)
    tree = slide.shapes._spTree
    tree.append(
        etree.fromstring(
            f'<p:grpSp xmlns:p="{P_NS}" xmlns:a="{A_NS}">'
            "<p:nvGrpSpPr><p:cNvPr id=\"9\" name=\"g\"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>"
            "<p:grpSpPr/>"
            '<p:sp><p:nvSpPr><p:cNvPr id="10" name="t"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr/><p:txBody><a:bodyPr/>"
            "<a:p><a:r><a:t>grouped caption</a:t></a:r></a:p>"
            "</p:txBody></p:sp></p:grpSp>"
        )
    )
    assert "grouped caption" in slide_lines(slide)


def test_an_equation_inside_alternate_content_becomes_latex(tmp_path: Path) -> None:
    """The headline case: this is where PowerPoint puts every equation."""
    presentation, slide = _deck(tmp_path)
    slide.shapes._spTree.append(
        etree.fromstring(
            f'<mc:AlternateContent xmlns:mc="{MC_NS}" xmlns:p="{P_NS}" '
            f'xmlns:a="{A_NS}" xmlns:m="{M_NS}">'
            '<mc:Choice Requires="a14">'
            '<p:sp><p:nvSpPr><p:cNvPr id="4" name="eq"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr/><p:txBody><a:bodyPr/><a:p>"
            "<a:r><a:t>written </a:t></a:r>"
            "<m:oMath><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e>"
            "<m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub></m:oMath>"
            "</a:p></p:txBody></p:sp>"
            "</mc:Choice>"
            "<mc:Fallback>"
            '<p:sp><p:nvSpPr><p:cNvPr id="5" name="img"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>PICTURE</a:t></a:r></a:p>"
            "</p:txBody></p:sp>"
            "</mc:Fallback></mc:AlternateContent>"
        )
    )
    lines = slide_lines(slide)
    assert "written $x_{0}$" in lines
    assert "PICTURE" not in lines, "the Fallback is a raster of the Choice, not extra content"


def test_the_fallback_is_used_when_there_is_no_choice(tmp_path: Path) -> None:
    """Some writers emit only a Fallback. Half a slide beats none of it."""
    presentation, slide = _deck(tmp_path)
    slide.shapes._spTree.append(
        etree.fromstring(
            f'<mc:AlternateContent xmlns:mc="{MC_NS}" xmlns:p="{P_NS}" xmlns:a="{A_NS}">'
            "<mc:Fallback>"
            '<p:sp><p:nvSpPr><p:cNvPr id="6" name="f"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            "<p:spPr/><p:txBody><a:bodyPr/><a:p><a:r><a:t>fallback text</a:t></a:r></a:p>"
            "</p:txBody></p:sp>"
            "</mc:Fallback></mc:AlternateContent>"
        )
    )
    assert "fallback text" in slide_lines(slide)


def test_a_table_becomes_a_markdown_table(tmp_path: Path) -> None:
    """A GraphicFrame has no text frame, so tables extracted as nothing at all."""
    presentation, slide = _deck(tmp_path)
    shape = slide.shapes.add_table(2, 2, 0, 0, 200, 100)
    table = shape.table
    table.cell(0, 0).text = "Base"
    table.cell(0, 1).text = "Digits"
    table.cell(1, 0).text = "Binary"
    table.cell(1, 1).text = "0,1"
    lines = slide_lines(slide)
    assert "| Base | Digits |" in lines
    assert "| --- | --- |" in lines
    assert "| Binary | 0,1 |" in lines


def test_speaker_notes_are_extracted_separately(tmp_path: Path) -> None:
    """Lecture decks often carry the actual explanation in the notes pane."""
    presentation, slide = _deck(tmp_path)
    slide.notes_slide.notes_text_frame.text = "Stress that delta depends on epsilon."
    assert speaker_notes(slide) == "Stress that delta depends on epsilon."


def test_notes_do_not_pick_up_the_slide_number_placeholder(tmp_path: Path) -> None:
    """A notes slide also holds a thumbnail and a slide-number placeholder."""
    presentation, slide = _deck(tmp_path)
    assert speaker_notes(slide) == ""


def test_extract_blocks_carries_slide_numbers_and_notes(tmp_path: Path) -> None:
    """End to end, through the public façade."""
    presentation, slide = _deck(tmp_path)
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "A limit is a value a function approaches."
    slide.notes_slide.notes_text_frame.text = "Mention the epsilon-delta game."
    blocks = extract_blocks(_save(presentation, tmp_path))

    assert len(blocks) == 1
    assert blocks[0].meta["slide"] == 1
    assert "A limit is a value a function approaches." in blocks[0].text
    assert "Mention the epsilon-delta game." in blocks[0].text
    assert blocks[0].meta["extraction"]["method"] == "text"


def test_an_empty_slide_produces_no_block(tmp_path: Path) -> None:
    """An all-image slide is silence, not an empty string in the corpus."""
    from pptx import Presentation

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    assert extract_blocks(_save(presentation, tmp_path)) == []
