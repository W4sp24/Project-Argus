"""Tests for Word extraction.

One block for a whole document means one citation for a whole document, and
a `docx` table is not a paragraph -- so a document whose content was a
comparison grid used to index as its prose only.
"""

from pathlib import Path

import pytest

from backend.rag.extract import extract_blocks

pytest.importorskip("docx", reason="the [rag] extra is not installed")


def _document():
    import docx

    return docx.Document()


def _save(document, tmp_path: Path) -> Path:
    path = tmp_path / "notes.docx"
    document.save(path)
    return path


def test_headings_split_the_document_into_sections(tmp_path: Path) -> None:
    document = _document()
    document.add_heading("Limits", level=1)
    document.add_paragraph("A limit is a value a function approaches.")
    document.add_heading("Continuity", level=1)
    document.add_paragraph("A function is continuous where its limit equals its value.")

    blocks = extract_blocks(_save(document, tmp_path))
    assert [block.meta.get("section") for block in blocks] == ["Limits", "Continuity"]
    assert "approaches" in blocks[0].text
    assert "approaches" not in blocks[1].text


def test_a_table_is_kept_as_markdown(tmp_path: Path) -> None:
    document = _document()
    document.add_heading("Bases", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Base"
    table.cell(0, 1).text = "Digits"
    table.cell(1, 0).text = "Binary"
    table.cell(1, 1).text = "0,1"

    text = extract_blocks(_save(document, tmp_path))[0].text
    assert "| Base | Digits |" in text
    assert "| --- | --- |" in text
    assert "| Binary | 0,1 |" in text


def test_a_table_stays_in_the_section_it_appears_in(tmp_path: Path) -> None:
    """document.tables is a separate list, so reading it in turn reorders."""
    document = _document()
    document.add_heading("First", level=1)
    first = document.add_table(rows=1, cols=1)
    first.cell(0, 0).text = "belongs to First"
    document.add_heading("Second", level=1)
    document.add_paragraph("no table here")

    blocks = {block.meta.get("section"): block.text for block in extract_blocks(
        _save(document, tmp_path)
    )}
    assert "belongs to First" in blocks["First"]
    assert "belongs to First" not in blocks["Second"]


def test_a_document_with_no_headings_is_one_block(tmp_path: Path) -> None:
    """The old behaviour, and the right answer for an unstructured document."""
    document = _document()
    document.add_paragraph("first")
    document.add_paragraph("second")

    blocks = extract_blocks(_save(document, tmp_path))
    assert len(blocks) == 1
    assert blocks[0].text == "first\nsecond"


def test_an_empty_document_yields_nothing(tmp_path: Path) -> None:
    assert extract_blocks(_save(_document(), tmp_path)) == []
